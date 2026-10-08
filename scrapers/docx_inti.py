"""
scrapers/docx_inti.py — Helper dokumen Word yang dipakai bersama.

Dua mesin dokumen hidup di proyek ini dan keduanya memakai helper di sini, jadi
tidak ada logika yang berganda:

  • `dokumen_isb.py` — MENGISI template berisi placeholder (proposal ISB).
  • `dokumen_buat.py` — MENYUSUN badan surat dari nol di atas kop NXTG, karena
    `dokumen/template NXTG.docx` hanya berisi kop surat (badannya satu paragraf
    kosong, tanpa satu pun placeholder).

Yang perlu diketahui sebelum menambah helper di sini:

  • Kop NXTG tidak bisa diandalkan untuk font. Gaya "Normal"-nya kosong,
    sehingga font & ukuran turun dari `docDefaults` = tema `minorHAnsi`
    (Calibri) 12pt dengan `spacing after=160 line=278`. Padahal kop suratnya
    Times New Roman. Karena itu `gaya_run()` dan `paragraf()` menulis font,
    ukuran, dan jarak paragraf SELALU eksplisit, bukan mengandalkan style.
  • python-docx tidak punya accessor untuk `tblW`/`tblBorders`/`tblCellMar`,
    dan Word merapikan ulang XML yang urutan anaknya salah. Urutan skema
    ditulis sekali di URUT_* dan dipakai lewat `anak()`.
  • Uang SELALU bilangan bulat. float untuk rupiah berarti angka yang tercetak
    di PDF tidak mau dijumlahkan.
"""
import re
import threading
from pathlib import Path

from docx.enum.table import WD_ALIGN_VERTICAL  # noqa: F401 — dipakai pemanggil
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips  # noqa: F401 — RGBColor dipakai pemanggil

BASE_DIR = Path(__file__).resolve().parent.parent
FOLDER_DOKUMEN = BASE_DIR / "dokumen"

BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
         "Agustus", "September", "Oktober", "November", "Desember"]
HARI = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]

_RE_SISA = re.compile(r"\[[^\]\n]{2,60}\]|\{\{[^}]+\}\}")
# Tanda klasifikasi sensitivity label Microsoft 365 ("[ OFFICIAL ]",
# "[ INTERNAL ]", ...) yang disisipkan Word ke header/footer — bukan placeholder.
# Placeholder template asli selalu huruf campuran, jadi tidak ikut terkecuali.
_RE_PENANDA = re.compile(r"^\[\s*[A-Z][A-Z0-9 \-:/]*\s*\]$")


def _sisa_placeholder(teks):
    """Placeholder yang masih tertinggal di `teks`, tanpa tanda klasifikasi Word.

    HANYA untuk dokumen hasil isi-template (dokumen_isb). JANGAN dipakai pada
    dokumen hasil susunan dokumen_buat: di sana tidak ada placeholder, tapi
    deskripsi barang dari pelanggan seperti 'Monitor LG 24" [refurbished]'
    akan terbaca sebagai placeholder dan membatalkan dokumen yang sah.
    """
    return {m.group() for m in _RE_SISA.finditer(teks)
            if not _RE_PENANDA.match(m.group())}


# ─── Format tanggal & nomor ───────────────────────────────────────────────────

def romawi(n):
    angka_romawi = [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    hasil = ""
    for nilai, huruf in angka_romawi:
        while n >= nilai:
            hasil += huruf
            n -= nilai
    return hasil


def tanggal_indonesia(tgl):
    return f"{tgl.day} {BULAN[tgl.month - 1]} {tgl.year}"


def hari_indonesia(tgl):
    return HARI[tgl.weekday()]


def bentuk_nomor(pola, urut, tgl):
    """Teks nomor surat dari pola pengaturan, mis. 001/ISB/SALES/26/IX/2026."""
    try:
        return pola.format(urut=urut, hari=tgl.day, bulan=tgl.month,
                           tahun=tgl.year, romawi=romawi(tgl.month))
    except (KeyError, IndexError, ValueError) as e:
        raise ValueError(f"Pola nomor surat tidak valid ({pola!r}): {e}") from None


def _nama_file(teks, batas=60):
    bersih = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", str(teks or "")).strip(" .")
    return re.sub(r"\s+", " ", bersih)[:batas].strip() or "Lead"


def folder_hari(akar, tgl):
    """Subfolder harian di bawah `akar`, mis. output/dokumen/2026-10-08."""
    return Path(akar) / tgl.strftime("%Y-%m-%d")


# ─── Uang ─────────────────────────────────────────────────────────────────────
#
# Rupiah SELALU int. float untuk uang berarti 0,1 + 0,2 != 0,3 dan angka cetak
# yang tidak mau dijumlahkan: subtotal di layar tidak sama dengan subtotal di
# PDF. Semua angka yang dicetak diambil dari bilangan bulat yang tersimpan,
# tidak pernah dihitung ulang saat render.

SATUAN_KATA = ("", "satu", "dua", "tiga", "empat", "lima", "enam", "tujuh",
               "delapan", "sembilan", "sepuluh", "sebelas")
SKALA = ((10 ** 6, "ribu"), (10 ** 9, "juta"), (10 ** 12, "miliar"),
         (10 ** 15, "triliun"))

_RE_BUKAN_UANG = re.compile(r"[^\d.,\-]")
_RE_DESIMAL = re.compile(r"[.,](\d{1,2})$")
# Bentuk yang diterima: '1250000', '1.250.000', '1,250,000', 'Rp 1.250.000',
# 'Rp. 1.250.000', '-1.250.000', dan akhiran gaya Indonesia 'Rp 1.000,-'.
_RE_UANG_SAH = re.compile(r"^-?\s*(?:rp\.?\s*)?-?[\d.,]+(?:[.,]-)?$", re.I)
# Pengelompokan ribuan yang sah. '1.2.3.4' ditolak: itu bukan angka yang
# ditulis orang, dan membiarkannya lewat akan menghasilkan 1234 secara diam-diam.
_RE_RIBUAN = re.compile(r"^\d+$|^\d{1,3}(?:[.,]\d{3})+$")


def bagi_bulat(a, b):
    """Pembagian bilangan bulat, dibulatkan ke atas pada 0,5. `a`, `b` >= 0."""
    a, b = int(a), int(b)
    if b <= 0:
        raise ValueError(f"Pembagi harus positif: {b}")
    if a < 0:
        return -((-a * 2 + b) // (2 * b))
    return (a * 2 + b) // (2 * b)


def ke_int(nilai, nama="nilai"):
    """
    'Rp 1.250.000' / '1,250,000' / 1250000.0 / '' → bilangan bulat.

    Bagian pecahan DITOLAK, tidak dipotong: impor Excel lewat pandas memberi
    float64 ('1250000.0'), dan memotong diam-diam angka seperti '1250000.75'
    akan membuat invoice tidak cocok dengan yang disepakati.

    Titik dan koma dianggap pemisah ribuan — konvensi Indonesia. Pemisah
    terakhir yang hanya diikuti 1-2 angka diperlakukan sebagai tanda desimal.
    """
    if nilai is None or nilai == "":
        return 0
    if isinstance(nilai, bool):
        raise ValueError(f"{nama} bukan angka: {nilai!r}")
    if isinstance(nilai, int):
        return nilai
    if isinstance(nilai, float):
        if not nilai.is_integer():
            raise ValueError(f"{nama} tidak boleh pecahan: {nilai!r}")
        return int(nilai)

    asli = str(nilai).strip()
    if not asli or asli == "-":
        return 0
    # Divalidasi SEBELUM dibersihkan: membuang karakter bukan angka lebih dulu
    # akan mengubah 'abc' menjadi '' dan terbaca sebagai nol, bukan galat.
    if not _RE_UANG_SAH.match(asli):
        raise ValueError(f"{nama} bukan angka: {nilai!r}")
    teks = _RE_BUKAN_UANG.sub("", asli).strip()
    if teks.endswith(",-") or teks.endswith(".-"):
        teks = teks[:-2]
    if not teks or teks == "-":
        return 0
    minus = teks.startswith("-")
    teks = teks.lstrip("-")
    # Pemisah terakhir yang hanya diikuti 1-2 angka = tanda desimal, bukan
    # ribuan (ribuan selalu berkelompok 3). Pecahan nol dibuang, selain itu
    # ditolak — bukan dipotong.
    m = _RE_DESIMAL.search(teks)
    if m:
        if int(m.group(1)) != 0:
            raise ValueError(f"{nama} tidak boleh pecahan: {nilai!r}")
        teks = teks[:m.start()]
    if not _RE_RIBUAN.match(teks):
        raise ValueError(f"{nama} bukan angka: {nilai!r}")
    digit = re.sub(r"[.,]", "", teks)
    return -int(digit) if minus else int(digit)


def angka(n):
    """1234567 → '1.234.567' (pemisah ribuan Indonesia)."""
    n = int(n)
    return ("-" if n < 0 else "") + f"{abs(n):,}".replace(",", ".")


def rupiah(n, awalan="Rp "):
    """1234567 → 'Rp 1.234.567'. Nol tetap dicetak 'Rp 0'."""
    return awalan + angka(n)


def _kata(n):
    if n < 12:
        return SATUAN_KATA[n]
    if n < 20:
        return _kata(n - 10) + " belas"
    if n < 100:
        return _kata(n // 10) + " puluh" + (" " + _kata(n % 10) if n % 10 else "")
    if n < 200:
        return "seratus" + (" " + _kata(n - 100) if n > 100 else "")
    if n < 1000:
        return _kata(n // 100) + " ratus" + (" " + _kata(n % 100) if n % 100 else "")
    if n < 2000:
        return "seribu" + (" " + _kata(n - 1000) if n > 1000 else "")
    for batas, nama in SKALA:
        if n < batas:
            satuan = batas // 1000
            return (_kata(n // satuan) + " " + nama
                    + (" " + _kata(n % satuan) if n % satuan else ""))
    raise ValueError(f"Angka terlalu besar untuk dieja: {n}")


def angka_ke_kata(n):
    """123450000 → 'seratus dua puluh tiga juta empat ratus lima puluh ribu'."""
    n = int(n)
    if n < 0:
        return "minus " + angka_ke_kata(-n)
    return "nol" if n == 0 else " ".join(_kata(n).split())


def rupiah_kata(n, akhiran=" rupiah"):
    """Untuk kwitansi & baris Total: '... empat ratus lima puluh ribu rupiah'."""
    return angka_ke_kata(n) + akhiran


# ─── Penelusuran & pengisian paragraf (dipakai dokumen_isb) ───────────────────

def _semua_paragraf(doc):
    """Paragraf badan, tabel (bersarang), header, dan footer."""
    def dari(wadah):
        for p in wadah.paragraphs:
            yield p
        for t in getattr(wadah, "tables", []):
            for baris in t.rows:
                for sel in baris.cells:
                    yield from dari(sel)
    yield from dari(doc)
    for s in doc.sections:
        for bagian in (s.header, s.first_page_header, s.even_page_header,
                       s.footer, s.first_page_footer, s.even_page_footer):
            if bagian is not None and not bagian.is_linked_to_previous:
                yield from dari(bagian)


def _hapus_sorotan(run):
    try:
        run.font.highlight_color = None
    except Exception:
        pass


def _isi_paragraf(p, ganti, nomor=None, pic=None):
    """Ganti placeholder di satu paragraf. Return True bila ada yang berubah."""
    runs = p.runs
    berubah = False

    # Baris "Nomor: [no. urut]/ISB/[bulan romawi]/2026" diganti utuh oleh nomor
    # dari pola pengaturan: sisa run di belakangnya dikosongkan.
    if nomor is not None and "[no. urut]" in p.text:
        for i, r in enumerate(runs):
            if "[no. urut]" in r.text:
                r.text = r.text.split("[no. urut]")[0] + nomor
                _hapus_sorotan(r)
                for r2 in runs[i + 1:]:
                    r2.text = ""
                return True

    for r in runs:
        teks = r.text
        for lama, baru in ganti.items():
            if lama in teks:
                teks = teks.replace(lama, baru)
        if teks != r.text:
            r.text = teks
            _hapus_sorotan(r)
            berubah = True

    if pic and p.text.strip() == "Bapak/Ibu Pimpinan":
        for r in runs:
            if "Bapak/Ibu Pimpinan" in r.text:
                r.text = r.text.replace("Bapak/Ibu Pimpinan", pic)
                berubah = True

    # Cadangan: placeholder yang terpecah ke beberapa run (terjadi bila template
    # diedit di Word). Teksnya digabung ke run pertama — formatnya ikut run itu.
    if runs and any(k in p.text for k in ganti):
        teks = p.text
        for lama, baru in ganti.items():
            teks = teks.replace(lama, baru)
        runs[0].text = teks
        _hapus_sorotan(runs[0])
        for r in runs[1:]:
            r.text = ""
        berubah = True
    return berubah


# ─── Konversi PDF lewat Microsoft Word ────────────────────────────────────────

_KUNCI_WORD = threading.Lock()


def ke_pdf(pasangan):
    """
    Konversi [(docx, pdf), ...] memakai satu instance Microsoft Word.

    Word COM tidak aman dipakai paralel, jadi dijaga satu kunci global. Kunci
    ini harus benar-benar SATU objek untuk seluruh proses: kalau fungsi ini
    disalin ke modul lain, dua kunci berarti dua Word berjalan bersamaan dan
    PDF-nya rusak. Karena itu modul lain mengimpor, tidak menyalin.

    Tiap file diputuskan sendiri: PDF lama yang sedang terbuka di pembaca PDF
    (terkunci) hanya menggagalkan file itu, bukan seluruh batch.

    Return {docx: None | pesan galat}. DOCX selalu tetap ada walau PDF gagal.
    """
    hasil = {asal: None for asal, _ in pasangan}
    if not pasangan:
        return hasil
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return {a: "pywin32 belum terpasang (pip install pywin32)" for a in hasil}
    with _KUNCI_WORD:
        pythoncom.CoInitialize()
        word = None
        try:
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0
            for asal, tujuan in pasangan:
                tujuan = Path(tujuan).resolve()
                try:
                    if tujuan.exists():
                        # Gagal di sini = file sedang dibuka aplikasi lain.
                        tujuan.unlink()
                    d = word.Documents.Open(str(Path(asal).resolve()), ReadOnly=True,
                                            AddToRecentFiles=False, Visible=False)
                    try:
                        d.SaveAs2(str(tujuan), FileFormat=17)  # wdFormatPDF
                    finally:
                        d.Close(False)
                except PermissionError:
                    hasil[asal] = (f"{tujuan.name} sedang dibuka aplikasi lain — tutup "
                                   f"dulu lalu buat ulang")
                except Exception as e:
                    hasil[asal] = f"{type(e).__name__}: {e}"
            return hasil
        except Exception as e:
            return {a: (v or f"Microsoft Word tidak bisa dijalankan ({type(e).__name__}: {e})")
                    for a, v in hasil.items()}
        finally:
            if word is not None:
                try:
                    word.Quit()
                except Exception:
                    pass
            pythoncom.CoUninitialize()


def ringkas_galat_pdf(hasil):
    """Satu kalimat peringatan dari hasil ke_pdf, atau None bila semua berhasil."""
    galat = [f"{Path(a).stem}: {g}" for a, g in hasil.items() if g]
    if not galat:
        return None
    return ("PDF gagal dibuat untuk " + "; ".join(galat[:5])
            + (f" (+{len(galat) - 5} lainnya)" if len(galat) > 5 else "")
            + ". File DOCX-nya tetap tersedia.")


# ─── Penyusun XML: urutan anak elemen menurut skema OOXML ─────────────────────
#
# Word menolak atau merapikan ulang XML yang urutan anaknya salah, dan
# python-docx TIDAK punya accessor untuk tblW/tblBorders/tblCellMar — hanya
# tblStyle, bidiVisual, jc, dan tblLayout. `tblPr.append(...)` menghasilkan anak
# ganda yang tidak urut. Urutannya ditulis sekali di sini.

URUT_TBLPR = ("w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
              "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
              "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
              "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
              "w:tblDescription", "w:tblPrChange")
URUT_TCPR = ("w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge",
             "w:tcBorders", "w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
             "w:tcFitText", "w:vAlign", "w:hideMark")
URUT_TRPR = ("w:cnfStyle", "w:divId", "w:gridBefore", "w:gridAfter", "w:wBefore",
             "w:wAfter", "w:cantSplit", "w:trHeight", "w:tblHeader",
             "w:tblCellSpacing", "w:jc", "w:hidden")
URUT_PPR = ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore",
            "w:framePr", "w:widowControl", "w:numPr", "w:suppressLineNumbers",
            "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
            "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
            "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
            "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
            "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
            "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
            "w:rPr", "w:sectPr", "w:pPrChange")
SISI = ("top", "left", "bottom", "right", "insideH", "insideV")


def el(tag, **atribut):
    """Elemen w:* baru dengan atribut w:*. `tipe=` dipakai untuk w:type."""
    e = OxmlElement(tag)
    for k, v in atribut.items():
        e.set(qn("w:" + ("type" if k == "tipe" else k)), str(v))
    return e


def anak(induk, tag, urut):
    """Anak `tag` yang sudah ada, atau yang baru disisipkan di posisi skema benar.

    `insert_element_before` adalah mekanisme python-docx sendiri (dipakai setiap
    ZeroOrOne), jadi urutannya pasti sah menurut skema.
    """
    ada = induk.find(qn(tag))
    if ada is not None:
        return ada
    baru = OxmlElement(tag)
    induk.insert_element_before(baru, *urut[urut.index(tag) + 1:])
    return baru


# ─── Font & paragraf ──────────────────────────────────────────────────────────

FONT = "Times New Roman"
PT_ISI = 11      # sama dengan baris alamat & telepon di kop surat (w:sz 22)
PT_JUDUL = 12
PT_KECIL = 9

RATA_KIRI = WD_ALIGN_PARAGRAPH.LEFT
RATA_TENGAH = WD_ALIGN_PARAGRAPH.CENTER
RATA_KANAN = WD_ALIGN_PARAGRAPH.RIGHT
RATA_PENUH = WD_ALIGN_PARAGRAPH.JUSTIFY

# Lebar kerja A4 dengan margin 1 inci: 11906 - 2*1440.
LEBAR_PAKAI = 9026


def gaya_run(run, pt=PT_ISI, tebal=False, miring=False, font=FONT, warna=None):
    """
    Pastikan satu run benar-benar `font` berukuran `pt`.

    Template NXTG tidak bisa diandalkan untuk ini: gaya "Normal"-nya kosong,
    jadi font & ukuran turun dari docDefaults = tema minorHAnsi (Calibri) 12pt.
    Keempat atribut rFonts ditulis semua — ascii/hAnsi untuk teks Latin,
    eastAsia supaya tanda baca tertentu tidak jatuh ke Calibri, cs untuk complex
    script (kop suratnya pun menulis w:cs). w:szCs ditulis manual karena
    python-docx hanya menulis w:sz.
    """
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)                       # rFonts = anak pertama CT_RPr
    for atr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(atr), font)
    run.font.size = Pt(pt)                      # menulis w:sz
    szCs = rPr.find(qn("w:szCs"))
    if szCs is None:
        szCs = OxmlElement("w:szCs")
        rPr.find(qn("w:sz")).addnext(szCs)      # szCs tepat setelah sz
    szCs.set(qn("w:val"), str(int(pt * 2)))
    run.bold = tebal or None
    run.italic = miring or None
    if warna is not None:
        run.font.color.rgb = warna
    return run


def paragraf(wadah, teks="", pt=PT_ISI, tebal=False, miring=False, rata=None,
             spasi_sebelum=0, spasi_sesudah=0, baris=1.0, kiri=0, gantung=0,
             jaga=False, gaya=None):
    """
    Satu paragraf dengan jarak & font yang ditulis eksplisit.

    `wadah` boleh Document, _Cell, _Header, atau _Footer. Jarak paragraf TIDAK
    diwariskan dengan benar dari template: docDefaults memberi
    `after=160 line=278` (1,15) yang membuat surat terlihat longgar, dan gaya
    "Normal" tidak menimpanya.

    `jaga=True` memakai keep_with_next/keep_together — perhatikan namanya:
    `keep_next` BUKAN atribut python-docx dan menuliskannya tidak menimbulkan
    galat sekaligus tidak berefek apa pun.
    """
    p = wadah.add_paragraph(style=gaya)
    pf = p.paragraph_format
    pf.space_before = Pt(spasi_sebelum)
    pf.space_after = Pt(spasi_sesudah)
    pf.line_spacing = baris
    if rata is not None:
        pf.alignment = rata
    if kiri:
        pf.left_indent = Twips(kiri)
    if gantung:
        pf.first_line_indent = Twips(-gantung)
        pf.tab_stops.add_tab_stop(Twips(kiri or gantung), WD_TAB_ALIGNMENT.LEFT)
    if jaga:
        pf.keep_with_next = True
        pf.keep_together = True
    if teks:
        gaya_run(p.add_run(teks), pt, tebal, miring)
    return p


def kosongkan_badan(doc):
    """
    Buang seluruh isi badan dokumen tapi pertahankan <w:sectPr>.

    sectPr adalah anak TERAKHIR <w:body> dan membawa ukuran A4, margin 1440,
    dan headerReference ke kop surat. Menghapus anak <w:body> satu per satu
    tanpa memeriksa sectPr menghancurkan tata letak halaman sekaligus kopnya.
    CT_Body.clear_content() sudah melakukan tepat itu, jadi dipakai langsung.

    Paragraf & tabel yang ditambahkan sesudah ini otomatis masuk SEBELUM
    sectPr, karena python-docx mendeklarasikan w:p/w:tbl dengan
    successors=('w:sectPr',).
    """
    doc.element.body.clear_content()
    return doc


def daftar_nomor(wadah, butir, kiri=567, pt=PT_ISI, spasi=3, rata=RATA_PENUH,
                 tanda=None):
    """
    Daftar bernomor tanpa numbering.xml.

    Template NXTG tidak punya word/numbering.xml maupun gaya "List Number"
    (python-docx melempar KeyError), dan Document.part.numbering_part masih
    NotImplementedError di python-docx 1.2.0 — menyuntik numPr berarti merakit
    sendiri part, content-type, dan relationship-nya. Nomor manual + indentasi
    menggantung memberi hasil visual yang sama, selalu stabil, dan nomornya
    ikut terbaca skrip verifikasi karena ada di dalam <w:t>.
    """
    hasil = []
    for i, teks in enumerate(butir, 1):
        label = tanda(i) if tanda else f"{i}."
        p = paragraf(wadah, pt=pt, rata=rata, kiri=kiri, gantung=kiri,
                     spasi_sesudah=spasi, baris=1.0)
        gaya_run(p.add_run(f"{label}\t{teks}"), pt)
        hasil.append(p)
    return hasil


# ─── Tabel ────────────────────────────────────────────────────────────────────

def tabel(doc, lebar, garis=True, sz=6, warna="000000", margin_sel=108, ind=0,
          baris=0):
    """
    Tabel lebar tetap dengan border yang ditulis eksplisit.

    `lebar` = daftar lebar kolom dalam twips; jumlahnya tidak boleh melebihi
    LEBAR_PAKAI. Border ditulis di tblPr walaupun gaya "Table Grid" ADA di
    template NXTG: kop suratnya sendiri memakai pola ini (tblStyle TableGrid
    plus tblBorders yang menimpanya), dan tabel item butuh ketebalan garis yang
    berbeda dari sz=4 seragam milik gaya itu. Hasilnya jadi tidak bergantung
    pada styles.xml sama sekali.

    `margin_sel=0` untuk tabel tata letak tanpa garis: dengan margin bawaan 108
    twips (dari gaya TableNormal) teks kolom pertama menjorok 0,075 inci
    dibanding paragraf biasa di atas dan di bawahnya — terlihat sebagai salah
    rata.
    """
    if sum(lebar) > LEBAR_PAKAI:
        raise ValueError(f"Lebar tabel {sum(lebar)} twip melebihi area cetak "
                         f"{LEBAR_PAKAI} twip")
    t = doc.add_table(rows=baris, cols=len(lebar))
    t.autofit = False                                   # → tblLayout fixed
    tblPr = t._tbl.tblPr

    tw = anak(tblPr, "w:tblW", URUT_TBLPR)              # menimpa w=0 type=auto
    tw.set(qn("w:w"), str(sum(lebar)))
    tw.set(qn("w:type"), "dxa")

    ti = anak(tblPr, "w:tblInd", URUT_TBLPR)
    ti.set(qn("w:w"), str(ind))
    ti.set(qn("w:type"), "dxa")

    tb = anak(tblPr, "w:tblBorders", URUT_TBLPR)
    for lama in list(tb):
        tb.remove(lama)
    for sisi in SISI:
        tb.append(el("w:" + sisi, val=("single" if garis else "none"),
                     sz=(sz if garis else 0), space=0, color=warna))

    tcm = anak(tblPr, "w:tblCellMar", URUT_TBLPR)
    for lama in list(tcm):
        tcm.remove(lama)
    for sisi in ("top", "left", "bottom", "right"):
        tcm.append(el("w:" + sisi,
                      w=(margin_sel if sisi in ("left", "right") else 0),
                      tipe="dxa"))

    # gridCol = dasar tata letak Word; tcW per sel yang menang di tblLayout
    # fixed. Keduanya harus diset, kalau hanya salah satu Word menghitung ulang
    # dan lebarnya tidak sesuai.
    for i, w in enumerate(lebar):
        t.columns[i].width = Twips(w)
    t._lebar = list(lebar)
    return t


def baris_tabel(t):
    """Baris baru dengan w:tcW terpasang di SETIAP sel."""
    r = t.add_row()
    for i, sel in enumerate(r.cells):
        sel.width = Twips(t._lebar[i])
    return r


def isi_sel(sel, teks, pt=PT_ISI, tebal=False, miring=False, rata=None,
            baris=1.0):
    """Tulis teks ke paragraf pertama sel, dengan font & jarak eksplisit."""
    p = sel.paragraphs[0]
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = baris
    if rata is not None:
        pf.alignment = rata
    if teks != "":
        gaya_run(p.add_run(str(teks)), pt, tebal, miring)
    return p


def sel_uang(sel, nilai, lebar, pt=PT_ISI, tebal=False, awalan="Rp"):
    """Sel uang: 'Rp' rata kiri, angkanya rata kanan lewat tab stop."""
    p = sel.paragraphs[0]
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pf.tab_stops.add_tab_stop(Twips(max(lebar - 216, 100)),
                              WD_TAB_ALIGNMENT.RIGHT)
    gaya_run(p.add_run(f"{awalan}\t{angka(nilai)}"), pt, tebal)
    return p


def warnai(sel, isian="D9D9D9"):
    """Latar satu sel (dipakai baris kepala tabel)."""
    shd = anak(sel._tc.get_or_add_tcPr(), "w:shd", URUT_TCPR)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), isian)


def tengah_vertikal(sel):
    anak(sel._tc.get_or_add_tcPr(), "w:vAlign", URUT_TCPR).set(qn("w:val"), "center")


def ulang_kepala(r):
    """Baris kepala diulang di tiap halaman dan tidak dipecah."""
    trPr = r._tr.get_or_add_trPr()
    anak(trPr, "w:cantSplit", URUT_TRPR)
    anak(trPr, "w:tblHeader", URUT_TRPR)


def jangan_pecah(r, tinggi_min=0):
    """Baris tidak boleh terpotong antar halaman; pindah utuh ke halaman baru."""
    trPr = r._tr.get_or_add_trPr()
    anak(trPr, "w:cantSplit", URUT_TRPR)
    if tinggi_min:
        th = anak(trPr, "w:trHeight", URUT_TRPR)
        th.set(qn("w:val"), str(int(tinggi_min)))
        th.set(qn("w:hRule"), "atLeast")


def jaga_bersama(t):
    """
    Seluruh tabel ini menempel ke blok sesudahnya (keepNext di tiap paragraf).

    Dipakai untuk blok kontak yang berdiri tepat di atas tanda tangan: tanpa
    ini, blok tanda tangan bisa pindah sendirian ke halaman berikutnya dan
    halaman itu hanya berisi "Hormat kami" — terlihat seperti salah cetak.
    Dengan keepNext, keduanya pindah bersama dan halaman barunya terlihat
    disengaja.
    """
    for baris in t.rows:
        jangan_pecah(baris)
        for sel in baris.cells:
            for p in sel.paragraphs:
                p.paragraph_format.keep_with_next = True
    return t


def garis_sel(sel, **sisi):
    """Border per sel, mis. garis_sel(c, top='double', bottom='none')."""
    tcB = anak(sel._tc.get_or_add_tcPr(), "w:tcBorders", URUT_TCPR)
    for nama, val in sisi.items():
        lama = tcB.find(qn("w:" + nama))
        if lama is not None:
            tcB.remove(lama)
        tcB.append(el("w:" + nama, val=val, sz=(0 if val == "none" else 6),
                      space=0, color="000000"))


def garis_bawah_paragraf(p, sz=6):
    """Garis di bawah satu paragraf — dipakai sebagai alas nama penanda tangan.

    pBdr bukan anak pertama pPr (pStyle/keepNext/... lebih dulu), jadi
    disisipkan lewat `anak()` memakai urutan skema, bukan insert(0).
    """
    pBdr = anak(p._p.get_or_add_pPr(), "w:pBdr", URUT_PPR)
    for lama in list(pBdr):
        pBdr.remove(lama)
    pBdr.append(el("w:bottom", val="single", sz=sz, space=1, color="000000"))
    return p


# ─── Footer & medan halaman ───────────────────────────────────────────────────

def medan(p, instr, pt=PT_KECIL, contoh="1"):
    """
    <w:fldSimple> mis. ' PAGE ' / ' NUMPAGES '. python-docx tidak punya API medan.

    Jangan memanggil p.add_run() lebih dulu — itu meninggalkan <w:r/> kosong.
    Word menghitung ulang medannya saat dokumen dibuka; `contoh` hanya yang
    terlihat di pembaca selain Word, dan ke_pdf membuka dokumen di Word.

    Run di dalam fldSimple TIDAK ikut `p.runs` (itu hanya anak langsung w:p),
    jadi fontnya diseragamkan di sini, bukan dengan menyapu p.runs sesudahnya.
    """
    from docx.text.run import Run

    f = el("w:fldSimple", instr=instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = contoh
    r.append(t)
    f.append(r)
    p._p.append(f)
    gaya_run(Run(r, p), pt)
    return f


def pasang_footer(doc, teks_kiri="", nomor_halaman=True, pt=PT_KECIL):
    """
    Footer berisi catatan kiri dan "Halaman X dari Y" di kanan.

    Template NXTG tidak punya footer part sama sekali, jadi invoice yang
    menjadi dua halaman tidak bernomor tanpa ini. Menyetel
    is_linked_to_previous = False membuat word/footer1.xml, relationship-nya,
    dan <w:footerReference> di sectPr yang sudah ada — headerReference dan
    pgMar tidak tersentuh.
    """
    s = doc.sections[0]
    s.footer.is_linked_to_previous = False
    p = s.footer.paragraphs[0]
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pf.tab_stops.clear_all()        # buang tab stop warisan gaya "Footer"
    pf.tab_stops.add_tab_stop(Twips(LEBAR_PAKAI), WD_TAB_ALIGNMENT.RIGHT)
    if teks_kiri:
        gaya_run(p.add_run(teks_kiri), pt)
    if nomor_halaman:
        gaya_run(p.add_run("\tHalaman "), pt)
        medan(p, " PAGE ", pt)
        gaya_run(p.add_run(" dari "), pt)
        medan(p, " NUMPAGES ", pt)
    return p


def lebar_pakai(doc):
    """Lebar area cetak section pertama dalam twips."""
    s = doc.sections[0]
    return s.page_width.twips - s.left_margin.twips - s.right_margin.twips
